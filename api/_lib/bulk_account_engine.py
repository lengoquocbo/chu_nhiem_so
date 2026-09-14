"""Tạo hàng loạt tài khoản từ danh sách học sinh đã import vào lớp."""
from __future__ import annotations

import base64
import io
import re
import unicodedata

from openpyxl import Workbook

from .admin_engine import create_student_login
from .supabase_client import get_supabase


def _ascii_token(value: str) -> str:
    value = value.replace("đ", "d").replace("Đ", "D")
    normalized = unicodedata.normalize("NFD", value)
    return "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn" and ch.isalnum()).lower()


def _normalized_name(value: str) -> str:
    return " ".join(value.strip().casefold().split())


def _valid_full_name(value: str) -> bool:
    return bool(value and re.fullmatch(r"[^\W\d_]+(?:[\s'-][^\W\d_]+)*", value, re.UNICODE))


def _credentials_base(full_name: str) -> tuple[str, str]:
    tokens = [_ascii_token(token) for token in full_name.split()]
    tokens = [token for token in tokens if token]
    if not tokens:
        raise ValueError("Không thể tạo tên đăng nhập từ họ tên")
    email_base = "".join(tokens[-2:])
    password = f"{tokens[-1]}123456@"
    return email_base, password


def _existing_auth_emails() -> set[str]:
    supabase = get_supabase()
    emails: set[str] = set()
    page = 1
    while True:
        users = supabase.auth.admin.list_users(page=page, per_page=1000)
        for user in users:
            if user.email:
                emails.add(user.email.strip().lower())
        if len(users) < 1000:
            break
        page += 1
    return emails


def _unique_email(base: str, used_emails: set[str]) -> str:
    suffix = 1
    while True:
        local_part = base if suffix == 1 else f"{base}{suffix}"
        email = f"{local_part}@gmail.com"
        if email not in used_emails:
            used_emails.add(email)
            return email
        suffix += 1


def _class_students(class_id: str) -> tuple[list[dict], set[str]]:
    supabase = get_supabase()
    enrollment_resp = (
        supabase.table("student_class_enrollments")
        .select("student_id, students(id, student_code, full_name)")
        .eq("class_id", class_id)
        .is_("left_at", None)
        .execute()
    )
    students = [row["students"] for row in (enrollment_resp.data or []) if row.get("students")]
    student_ids = [student["id"] for student in students]
    linked: set[str] = set()
    if student_ids:
        link_resp = supabase.table("student_user_link").select("student_id").in_("student_id", student_ids).execute()
        linked = {row["student_id"] for row in (link_resp.data or [])}
    return students, linked


def preview_class_accounts(class_id: str) -> dict:
    """Tạo bản xem trước từ danh sách học sinh đã được import vào lớp."""
    students, linked_student_ids = _class_students(class_id)
    students.sort(
        key=lambda student: (
            str(student.get("student_code") or ""),
            _normalized_name(student.get("full_name") or ""),
            student["id"],
        )
    )
    used_emails = _existing_auth_emails()
    preview_rows: list[dict] = []

    for index, student in enumerate(students, start=1):
        full_name = " ".join(str(student.get("full_name") or "").strip().split())
        errors: list[str] = []
        if not full_name:
            errors.append("Thiếu họ tên")
        elif not _valid_full_name(full_name):
            errors.append("Họ tên sai định dạng")
        if student["id"] in linked_student_ids:
            errors.append("Học sinh đã có tài khoản")

        email = ""
        password = ""
        if not errors:
            try:
                email_base, password = _credentials_base(full_name)
                email = _unique_email(email_base, used_emails)
            except ValueError as exc:
                errors.append(str(exc))

        preview_rows.append(
            {
                "row_index": index,
                "student_id": student["id"],
                "student_code": student.get("student_code"),
                "full_name": full_name,
                "email": email,
                "password": password,
                "errors": errors,
                "is_valid": not errors,
            }
        )

    return {
        "total": len(preview_rows),
        "valid_count": sum(1 for row in preview_rows if row["is_valid"]),
        "invalid_count": sum(1 for row in preview_rows if not row["is_valid"]),
        "rows": preview_rows,
    }


def _excel_safe(value: str) -> str:
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _result_workbook(created: list[dict], failed: list[dict]) -> str:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tài khoản đã tạo"
    sheet.append(["STT", "Họ và tên", "Tên đăng nhập", "Mật khẩu mặc định", "Yêu cầu"])
    for index, item in enumerate(created, start=1):
        sheet.append([index, _excel_safe(item["full_name"]), item["email"], item["password"], "Đổi mật khẩu lần đầu"])
    for width, column in zip((8, 30, 34, 24, 26), "ABCDE"):
        sheet.column_dimensions[column].width = width

    if failed:
        failed_sheet = workbook.create_sheet("Không tạo được")
        failed_sheet.append(["STT", "Họ và tên", "Lý do"])
        for index, item in enumerate(failed, start=1):
            failed_sheet.append([index, _excel_safe(item.get("full_name") or ""), _excel_safe(item["reason"])])
        failed_sheet.column_dimensions["A"].width = 8
        failed_sheet.column_dimensions["B"].width = 30
        failed_sheet.column_dimensions["C"].width = 60

    output = io.BytesIO()
    workbook.save(output)
    return base64.b64encode(output.getvalue()).decode("ascii")


def export_student_accounts_xlsx(class_id: str) -> dict:
    """Xuất danh sách tài khoản học sinh hiện tại của lớp ra file xlsx (base64)."""
    from .admin_engine import list_class_student_accounts

    accounts = list_class_student_accounts(class_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Danh sách tài khoản"
    sheet.append(["STT", "Mã HS", "Họ và tên", "Email đăng nhập", "Trạng thái"])
    for index, item in enumerate(accounts, start=1):
        account = item.get("account")
        email = account["email"] if account else ""
        status = "Đã có tài khoản" if account else "Chưa có tài khoản"
        sheet.append([
            index,
            _excel_safe(str(item.get("student_code") or "")),
            _excel_safe(item["full_name"]),
            _excel_safe(email),
            status,
        ])
    for width, column in zip((8, 14, 30, 34, 20), "ABCDE"):
        sheet.column_dimensions[column].width = width

    output = io.BytesIO()
    workbook.save(output)
    return {
        "download_base64": base64.b64encode(output.getvalue()).decode("ascii"),
        "download_filename": "danh-sach-tai-khoan-hoc-sinh.xlsx",
    }


def create_bulk_accounts(class_id: str, selected_rows: list[dict], actor_id: str) -> dict:
    students, linked_student_ids = _class_students(class_id)
    students_by_id = {student["id"]: student for student in students}
    used_emails = _existing_auth_emails()
    created: list[dict] = []
    failed: list[dict] = []
    seen: set[str] = set()

    for selected in selected_rows:
        student_id = str(selected.get("student_id") or "")
        student = students_by_id.get(student_id)
        if not student:
            failed.append({"student_id": student_id, "full_name": selected.get("full_name") or "", "reason": "Học sinh không thuộc lớp đang phụ trách"})
            continue
        if student_id in seen:
            failed.append({"student_id": student_id, "full_name": student["full_name"], "reason": "Học sinh bị chọn trùng"})
            continue
        seen.add(student_id)
        if student_id in linked_student_ids:
            failed.append({"student_id": student_id, "full_name": student["full_name"], "reason": "Học sinh đã có tài khoản"})
            continue

        try:
            email_base, password = _credentials_base(student["full_name"])
            email = _unique_email(email_base, used_emails)
            account = create_student_login(student_id, email, actor_id, password, class_id)
            created.append({"student_id": student_id, "full_name": student["full_name"], "email": account["email"], "password": password})
            linked_student_ids.add(student_id)
        except Exception as exc:  # noqa: BLE001
            detail = getattr(exc, "detail", str(exc))
            failed.append({"student_id": student_id, "full_name": student["full_name"], "reason": str(detail)})

    return {
        "created_count": len(created),
        "failed_count": len(failed),
        "created": created,
        "failed": failed,
        "download_base64": _result_workbook(created, failed),
        "download_filename": "tai-khoan-hoc-sinh.xlsx",
    }
