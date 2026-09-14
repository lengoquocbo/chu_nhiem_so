"""
Import danh sách học sinh bằng Excel.

Luồng: Upload file -> validate() trả về preview + lỗi (KHÔNG ghi DB) ->
frontend hiển thị cho GVCN xem lại -> confirm() mới thực sự tạo học sinh.

Cột bắt buộc trong file (không bắt buộc đúng thứ tự):
  Họ và tên | Giới tính | Ngày tháng năm sinh | Số điện thoại phụ huynh |
  Email phụ huynh | Địa chỉ
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import load_workbook

from .supabase_client import get_supabase

EXPECTED_HEADERS = [
    "họ và tên",
    "giới tính",
    "ngày tháng năm sinh",
    "số điện thoại phụ huynh",
    "email phụ huynh",
    "địa chỉ",
]
PHONE_REGEX = re.compile(r"^(0|\+84)\d{9,10}$")
EMAIL_REGEX = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


@dataclass
class ImportRow:
    row_index: int
    full_name: str = ""
    gender: str | None = None
    dob: str | None = None  # ISO date string
    parent_phone: str | None = None
    parent_email: str | None = None
    address: str | None = None
    errors: list[str] = field(default_factory=list)

    @property
    def is_valid(self) -> bool:
        return len(self.errors) == 0


def _normalize_gender(raw: str | None) -> str | None:
    if not raw:
        return None
    v = raw.strip().lower()
    if v in ("nam", "male", "m"):
        return "male"
    if v in ("nữ", "nu", "female", "f"):
        return "female"
    return "other"


def _parse_date(raw) -> tuple[str | None, str | None]:
    """Trả về (iso_date, error_message)."""
    if raw in (None, ""):
        return None, "Thiếu ngày sinh"
    if isinstance(raw, datetime):
        return raw.date().isoformat(), None
    if isinstance(raw, date):
        return raw.isoformat(), None
    text = str(raw).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat(), None
        except ValueError:
            continue
    return None, f"Ngày sinh không hợp lệ: '{text}'"


def parse_and_validate(file_bytes: bytes, class_id: str) -> list[ImportRow]:
    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"Không đọc được file Excel: {exc}") from exc

    ws = wb.active
    rows_iter = list(ws.iter_rows(values_only=True))
    if not rows_iter:
        raise HTTPException(status_code=422, detail="File rỗng.")

    header = [str(h).strip().lower() if h else "" for h in rows_iter[0]]
    for expected in EXPECTED_HEADERS:
        if expected not in header:
            raise HTTPException(
                status_code=422,
                detail=f"Thiếu cột bắt buộc '{expected}'. Cột hiện có: {rows_iter[0]}",
            )
    col_idx = {name: header.index(name) for name in EXPECTED_HEADERS}

    supabase = get_supabase()
    existing_resp = (
        supabase.table("student_class_enrollments")
        .select("students(full_name, date_of_birth)")
        .eq("class_id", class_id)
        .is_("left_at", None)
        .execute()
    )
    existing_pairs = {
        (r["students"]["full_name"].strip().lower(), r["students"]["date_of_birth"])
        for r in (existing_resp.data or [])
        if r.get("students")
    }

    results: list[ImportRow] = []
    seen_in_file: set[tuple[str, str | None]] = set()

    for i, raw_row in enumerate(rows_iter[1:], start=2):
        if raw_row is None or all(v in (None, "") for v in raw_row):
            continue  # bỏ qua dòng trống

        row = ImportRow(row_index=i)
        row.full_name = str(raw_row[col_idx["họ và tên"]] or "").strip()
        row.gender = _normalize_gender(raw_row[col_idx["giới tính"]])
        dob_iso, dob_err = _parse_date(raw_row[col_idx["ngày tháng năm sinh"]])
        row.dob = dob_iso
        row.parent_phone = str(raw_row[col_idx["số điện thoại phụ huynh"]] or "").strip() or None
        row.parent_email = str(raw_row[col_idx["email phụ huynh"]] or "").strip().lower() or None
        row.address = str(raw_row[col_idx["địa chỉ"]] or "").strip() or None

        if not row.full_name:
            row.errors.append("Thiếu họ và tên")
        if dob_err:
            row.errors.append(dob_err)
        if row.parent_phone and not PHONE_REGEX.match(row.parent_phone):
            row.errors.append(f"Số điện thoại không hợp lệ: '{row.parent_phone}'")
        if row.parent_email and not EMAIL_REGEX.match(row.parent_email):
            row.errors.append(f"Email phụ huynh không hợp lệ: '{row.parent_email}'")

        dedup_key = (row.full_name.lower(), row.dob)
        if dedup_key in seen_in_file:
            row.errors.append("Trùng với 1 dòng khác trong cùng file (tên + ngày sinh)")
        seen_in_file.add(dedup_key)

        if dedup_key in existing_pairs:
            row.errors.append("Học sinh này (tên + ngày sinh) đã tồn tại trong lớp")

        results.append(row)

    if not results:
        raise HTTPException(status_code=422, detail="Không có dòng dữ liệu hợp lệ nào trong file.")

    return results


def confirm_import(rows: list[dict], class_id: str, school_year_id: str, actor_id: str) -> dict:
    """
    rows: danh sách dict đã qua validate ở frontend (chỉ những dòng is_valid=True
    được phép gửi lên confirm). Server VẪN re-validate cơ bản (không tin frontend).
    """
    supabase = get_supabase()
    created, failed = [], []

    for row in rows:
        full_name = (row.get("full_name") or "").strip()
        dob = row.get("dob")
        parent_phone = (row.get("parent_phone") or "").strip() or None
        parent_email = (row.get("parent_email") or "").strip().lower() or None
        if not full_name or not dob:
            failed.append({"row": row, "reason": "Thiếu họ tên hoặc ngày sinh"})
            continue
        if parent_phone and not PHONE_REGEX.match(parent_phone):
            failed.append({"row": row, "reason": "Số điện thoại phụ huynh không hợp lệ"})
            continue
        if parent_email and not EMAIL_REGEX.match(parent_email):
            failed.append({"row": row, "reason": "Email phụ huynh không hợp lệ"})
            continue

        student = None
        try:
            student_resp = (
                supabase.table("students")
                .insert(
                    {
                        "full_name": full_name,
                        "gender": row.get("gender"),
                        "date_of_birth": dob,
                        "parent_phone": parent_phone,
                        "parent_email": parent_email,
                        "address": row.get("address"),
                    }
                )
                .execute()
            )
            student = student_resp.data[0]

            supabase.table("student_class_enrollments").insert(
                {
                    "student_id": student["id"],
                    "class_id": class_id,
                    "school_year_id": school_year_id,
                    "created_by": actor_id,
                }
            ).execute()

            created.append(student)
        except Exception as exc:  # noqa: BLE001
            if student:
                try:
                    supabase.table("students").delete().eq("id", student["id"]).execute()
                except Exception:  # noqa: BLE001
                    pass
            failed.append({"row": row, "reason": str(exc)})

    supabase.table("audit_logs").insert(
        {
            "actor_id": actor_id,
            "action": "IMPORT",
            "entity": "students",
            "entity_id": class_id,
            "before_data": None,
            "after_data": {"created_count": len(created), "failed_count": len(failed)},
            "reason": f"Import Excel học sinh vào lớp {class_id}",
        }
    ).execute()

    return {"created": created, "failed": failed}
